"""Certify the global optimum for the 100 frozen ReCellular scenarios.

This script reads Simulation!B2:B101 from the existing workbook. It does not
generate demand, modify the workbook, or re-use the independent 10,000-sample
validation data.
"""

import argparse
import csv
import hashlib
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
import highspy


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OFFICIAL_WORKBOOK = PROJECT_ROOT / "data" / "ReCellular_Inventory_Planning.xlsx"
DEFAULT_RESULT = PROJECT_ROOT / "results" / "global_optimality_validation_evidence.csv"
DEFAULT_LOG = PROJECT_ROOT / "results" / "global_optimality_validation_highs.log"
TOLERANCE = 1e-6


def excel_logic(
    demand: np.ndarray,
    q: np.ndarray,
    acquisition: np.ndarray,
    remanufacturing: np.ndarray,
    price: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Apply the workbook's High -> Medium -> Low allocation and profit logic."""
    high = np.minimum(demand, q[0])
    remaining = np.maximum(demand - high, 0)
    medium = np.minimum(remaining, q[1])
    remaining = np.maximum(remaining - medium, 0)
    low = np.minimum(remaining, q[2])
    sold = np.column_stack([high, medium, low])
    profit = price * sold.sum(axis=1) - acquisition @ q - sold @ remanufacturing
    return sold, profit


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=OFFICIAL_WORKBOOK,
        help=f"Workbook path (default: {OFFICIAL_WORKBOOK})",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_RESULT,
        help=f"Solver evidence CSV (default: {DEFAULT_RESULT})",
    )
    parser.add_argument(
        "--log",
        type=Path,
        default=DEFAULT_LOG,
        help=f"Complete HiGHS log (default: {DEFAULT_LOG})",
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_demand(demand: np.ndarray) -> str:
    return hashlib.sha256(np.asarray(demand, dtype="<i8").tobytes()).hexdigest()


def portable_path(path: Path) -> str:
    try:
        return str(path.relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path)


def main() -> None:
    args = parse_args()
    workbook = args.workbook.expanduser().resolve()
    result_path = args.output.expanduser().resolve()
    log_path = args.log.expanduser().resolve()
    wb_values = load_workbook(workbook, data_only=True, read_only=True)
    inputs = wb_values["Inputs"]
    simulation = wb_values["Simulation"]

    demand = np.array([simulation.cell(row, 2).value for row in range(2, 102)], dtype=int)
    acquisition = np.array([inputs["B2"].value, inputs["C2"].value, inputs["D2"].value], dtype=float)
    remanufacturing = np.array([inputs["B3"].value, inputs["C3"].value, inputs["D3"].value], dtype=float)
    prices = np.array([inputs["B4"].value, inputs["C4"].value, inputs["D4"].value], dtype=float)

    np.testing.assert_array_equal(acquisition, [50, 35, 5])
    np.testing.assert_array_equal(remanufacturing, [30, 50, 85])
    np.testing.assert_array_equal(prices, [120, 120, 120])
    price = float(prices[0])
    n_scenarios = len(demand)

    model = highspy.Highs()
    log_path.parent.mkdir(parents=True, exist_ok=True)
    log_path.write_text("", encoding="utf-8")
    model.setOptionValue("output_flag", True)
    model.setOptionValue("log_to_console", False)
    model.setOptionValue("log_file", str(log_path))
    model.setOptionValue("mip_rel_gap", 0.0)

    # Buying more than maximum observed demand can only add acquisition cost.
    q = model.addVariables(
        3,
        lb=0,
        ub=int(demand.max()),
        obj=(-acquisition).tolist(),
        type=highspy.HighsVarType.kInteger,
        name=["Q_H", "Q_M", "Q_L"],
        out_array=True,
    )

    r = {}
    for s in range(n_scenarios):
        for g, grade in enumerate(("H", "M", "L")):
            r[s, g] = model.addVariable(
                lb=0,
                ub=int(demand[s]),
                obj=(price - remanufacturing[g]) / n_scenarios,
                type=highspy.HighsVarType.kInteger,
                name=f"R_{grade}_{s + 1}",
            )
            model.addConstr(r[s, g] <= q[g], name=f"inventory_{grade}_{s + 1}")
        model.addConstr(sum(r[s, g] for g in range(3)) <= int(demand[s]), name=f"demand_{s + 1}")

    model.setMaximize()
    model.run()

    status = model.modelStatusToString(model.getModelStatus())
    info = model.getInfo()
    if status != "Optimal":
        raise RuntimeError(f"HiGHS did not prove optimality: {status}; gap={info.mip_gap}")

    solution = model.getSolution()
    q_raw = np.asarray(solution.col_value[:3], dtype=float)
    if np.any(q_raw < -TOLERANCE) or not np.allclose(q_raw, np.rint(q_raw), atol=TOLERANCE, rtol=0):
        raise AssertionError(f"Purchase quantities are not nonnegative integers: {q_raw.tolist()}")
    q_star = np.rint(q_raw).astype(int)

    optimized_sales = np.array(
        [[solution.col_value[3 + 3 * s + g] for g in range(3)] for s in range(n_scenarios)]
    )
    optimized_scenario_profit = (
        price * optimized_sales.sum(axis=1)
        - acquisition @ q_star
        - optimized_sales @ remanufacturing
    )

    excel_sales, excel_profit = excel_logic(demand, q_star, acquisition, remanufacturing, price)
    if not np.array_equal(np.rint(optimized_sales).astype(int), excel_sales.astype(int)):
        raise AssertionError("MILP recourse allocation differs from the Excel High->Medium->Low logic.")

    optimal_profit = float(excel_profit.mean())
    objective_difference = optimal_profit - float(info.objective_function_value)
    scenario_difference = float(np.max(np.abs(optimized_scenario_profit - excel_profit)))
    if abs(objective_difference) > TOLERANCE or scenario_difference > TOLERANCE:
        raise AssertionError(
            "MILP objective does not match direct profit recalculation: "
            f"objective difference={objective_difference}, scenario difference={scenario_difference}"
        )

    result = {
        "Workbook": portable_path(workbook),
        "Workbook SHA-256": sha256_file(workbook),
        "Demand SHA-256": sha256_demand(demand),
        "Scenario Count": n_scenarios,
        "HiGHS Version": model.version(),
        "Configured MIP Relative Gap": 0.0,
        "Log To Console": False,
        "Log File": portable_path(log_path),
        "Solver Status": status,
        "Optimal High": int(q_star[0]),
        "Optimal Medium": int(q_star[1]),
        "Optimal Low": int(q_star[2]),
        "Objective": float(info.objective_function_value),
        "Average Profit": optimal_profit,
        "Best Bound": float(info.mip_dual_bound),
        "MIP Gap": float(info.mip_gap),
        "Profit Recalculation Difference": objective_difference,
        "Maximum Scenario Profit Difference": scenario_difference,
    }
    result_path.parent.mkdir(parents=True, exist_ok=True)
    with result_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=result.keys())
        writer.writeheader()
        writer.writerow(result)

    print(f"Solver status: {status}")
    print(f"HiGHS version: {model.version()}")
    print(f"MIP gap: {info.mip_gap:.12g}")
    print(f"Best bound: {info.mip_dual_bound:.12g}")
    print(f"Optimal Q: H={q_star[0]}, M={q_star[1]}, L={q_star[2]}")
    print(f"MILP average profit: {info.objective_function_value:.2f}")
    print(f"Excel-logic average profit: {optimal_profit:.2f}")
    print(f"Profit recalculation difference: {objective_difference:.12g}")
    print(f"Maximum scenario profit difference: {scenario_difference:.12g}")
    print(f"Result saved: {result_path}")
    print(f"HiGHS log saved: {log_path}")


if __name__ == "__main__":
    main()
